from __future__ import annotations

import json
from datetime import datetime
from pathlib import Path
from typing import Any

import pandas as pd


BASE_DIR = Path(__file__).resolve().parents[1]
OUTPUT_DIR = BASE_DIR / "runtime_outputs"
STATE_DIR = BASE_DIR / "state"
EXCEL_PATH = OUTPUT_DIR / "DAILY_MARKET_RECORD.xlsx"
HISTORY_PATH = STATE_DIR / "daily_order_history.csv"
REPORT_PATH = OUTPUT_DIR / "LATEST_EXCEL_RECORD_REPORT.json"


def now_text() -> str:
    return datetime.now().replace(microsecond=0).isoformat()


def read_json(path: Path, default: Any) -> Any:
    if not path.exists():
        return default
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except Exception:
        return default


def read_csv(path: Path) -> pd.DataFrame:
    if not path.exists():
        return pd.DataFrame()
    try:
        return pd.read_csv(path, low_memory=False)
    except Exception:
        return pd.DataFrame()


def write_json(path: Path, data: dict[str, Any]) -> None:
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")


def flatten_summary(summary: dict[str, Any]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    generated_at = summary.get("generated_at", "")
    for market in ["TW", "US", "US_SCAN"]:
        block = summary.get(market)
        if not isinstance(block, dict):
            continue
        row: dict[str, Any] = {"generated_at": generated_at, "market": market}
        for key, value in block.items():
            if isinstance(value, (list, dict)):
                continue
            row[key] = value
        rows.append(row)
    return pd.DataFrame(rows)


def flatten_price_update(report: dict[str, Any]) -> pd.DataFrame:
    rows: list[dict[str, Any]] = []
    for item in report.get("files", []) if isinstance(report, dict) else []:
        before = item.get("before") or {}
        after = item.get("after") or {}
        rows.append(
            {
                "started_at": report.get("started_at"),
                "finished_at": report.get("finished_at"),
                "label": item.get("label"),
                "market": item.get("market"),
                "status": item.get("status"),
                "new_rows": item.get("new_rows"),
                "attempted_symbols": item.get("attempted_symbols"),
                "updated_symbols": item.get("updated_symbols"),
                "error_count": item.get("error_count"),
                "before_rows": before.get("rows"),
                "before_symbols": before.get("symbols"),
                "before_max_date": before.get("max_date"),
                "after_rows": after.get("rows"),
                "after_symbols": after.get("symbols"),
                "after_max_date": after.get("max_date"),
                "path": item.get("path") or after.get("path") or before.get("path"),
            }
        )
    if not rows and isinstance(report, dict):
        rows.append(
            {
                "started_at": report.get("started_at"),
                "finished_at": report.get("finished_at"),
                "status": report.get("status"),
                "new_rows_total": report.get("new_rows_total"),
                "source_rate_limited": report.get("source_rate_limited"),
            }
        )
    return pd.DataFrame(rows)


def flatten_auto_rebalance(report: dict[str, Any]) -> pd.DataFrame:
    if not isinstance(report, dict) or not report:
        return pd.DataFrame()
    rows: list[dict[str, Any]] = []
    base = {
        "started_at": report.get("started_at"),
        "finished_at": report.get("finished_at"),
        "status": report.get("status"),
        "due_markets": ",".join(report.get("due_markets") or []),
        "full_markets": ",".join(report.get("full_markets") or []),
        "decision_text": report.get("decision_text"),
    }
    checks = report.get("due_checks") or []
    if checks:
        for check in checks:
            row = dict(base)
            row.update(
                {
                    "market": check.get("market"),
                    "due": check.get("due"),
                    "days_since_rebalance": check.get("estimated_weekday_trading_days_since_rebalance"),
                    "rebalance_step": check.get("rebalance_step_trading_days"),
                    "quote_date": check.get("current_quote_latest_date"),
                    "component_rebalance_date": check.get("component_rebalance_date"),
                }
            )
            rows.append(row)
    else:
        rows.append(base)
    return pd.DataFrame(rows)


def target_positions() -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    for market in ["TW", "US"]:
        df = read_csv(OUTPUT_DIR / market / "latest_target_position.csv")
        if df.empty:
            continue
        df = df.copy()
        df["market"] = market
        df = df[["market"] + [col for col in df.columns if col != "market"]]
        frames.append(df)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def market_states() -> pd.DataFrame:
    frames: list[pd.DataFrame] = []
    for market in ["TW", "US", "US_SCAN"]:
        df = read_csv(OUTPUT_DIR / market / "latest_market_state.csv")
        if df.empty:
            continue
        df = df.copy()
        df["market"] = market
        df = df[["market"] + [col for col in df.columns if col != "market"]]
        frames.append(df)
    return pd.concat(frames, ignore_index=True) if frames else pd.DataFrame()


def append_order_history(order_plan: pd.DataFrame, snapshot_at: str, exported_at: str) -> pd.DataFrame:
    STATE_DIR.mkdir(parents=True, exist_ok=True)
    if order_plan.empty:
        history = read_csv(HISTORY_PATH)
        return history

    keep_cols = [
        "market",
        "symbol",
        "name",
        "theme",
        "target_weight",
        "price_date",
        "target_amount",
        "target_shares",
        "current_shares",
        "diff_shares",
        "action",
        "entry_status",
        "entry_gap_pct",
        "entry_current_price",
        "entry_signal_date",
        "entry_signal_close",
        "buy_lower_price",
        "buy_upper_price",
        "tracking_base_date",
        "tracking_base_price",
        "tracking_return_pct",
    ]
    new_rows = order_plan.copy()
    for col in keep_cols:
        if col not in new_rows.columns:
            new_rows[col] = ""
    new_rows = new_rows[keep_cols].copy()
    new_rows.insert(0, "snapshot_at", snapshot_at)
    new_rows.insert(1, "exported_at", exported_at)

    old_rows = read_csv(HISTORY_PATH)
    history = pd.concat([old_rows, new_rows], ignore_index=True) if not old_rows.empty else new_rows
    if {"snapshot_at", "market", "symbol"}.issubset(history.columns):
        history = history.drop_duplicates(subset=["snapshot_at", "market", "symbol"], keep="last")
    history.to_csv(HISTORY_PATH, index=False, encoding="utf-8-sig")
    return history


def fit_workbook(path: Path) -> None:
    try:
        from openpyxl import load_workbook
    except Exception:
        return

    wb = load_workbook(path)
    for ws in wb.worksheets:
        ws.freeze_panes = "A2"
        if ws.max_row >= 1 and ws.max_column >= 1:
            ws.auto_filter.ref = ws.dimensions
        for column_cells in ws.columns:
            letter = column_cells[0].column_letter
            max_len = 8
            for cell in column_cells[:200]:
                value = "" if cell.value is None else str(cell.value)
                max_len = max(max_len, min(len(value), 42))
            ws.column_dimensions[letter].width = max_len + 2
    wb.save(path)


def export_excel_record() -> dict[str, Any]:
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    STATE_DIR.mkdir(parents=True, exist_ok=True)

    exported_at = now_text()
    summary = read_json(OUTPUT_DIR / "LATEST_DAILY_MARKET_POOL_SUMMARY.json", {})
    price_update = read_json(OUTPUT_DIR / "LATEST_PRICE_UPDATE_REPORT.json", {})
    target_refresh = read_json(OUTPUT_DIR / "LATEST_TARGET_PRICE_REFRESH_REPORT.json", {})
    auto_rebalance = read_json(OUTPUT_DIR / "LATEST_AUTO_REBALANCE_REPORT.json", {})
    snapshot_at = str(summary.get("generated_at") or exported_at)

    order_plan = read_csv(OUTPUT_DIR / "LATEST_ORDER_PLAN.csv")
    order_history = append_order_history(order_plan, snapshot_at=snapshot_at, exported_at=exported_at)

    sheets: dict[str, pd.DataFrame] = {
        "Today_Order_Plan": order_plan,
        "Order_History": order_history,
        "Market_State": market_states(),
        "Target_Position": target_positions(),
        "Current_Prices": read_csv(STATE_DIR / "current_prices.csv"),
        "Current_Holdings": read_csv(STATE_DIR / "current_holdings.csv"),
        "Update_Report": flatten_price_update(price_update),
        "Auto_Rebalance": flatten_auto_rebalance(auto_rebalance),
        "Summary": flatten_summary(summary),
        "Target_Price_Refresh": pd.DataFrame([target_refresh]) if isinstance(target_refresh, dict) else pd.DataFrame(),
    }

    with pd.ExcelWriter(EXCEL_PATH, engine="openpyxl") as writer:
        for name, df in sheets.items():
            sheet = name[:31]
            if df.empty:
                pd.DataFrame([{"note": "no data"}]).to_excel(writer, sheet_name=sheet, index=False)
            else:
                df.to_excel(writer, sheet_name=sheet, index=False)
    fit_workbook(EXCEL_PATH)

    report = {
        "ok": True,
        "exported_at": exported_at,
        "snapshot_at": snapshot_at,
        "excel_path": str(EXCEL_PATH),
        "history_path": str(HISTORY_PATH),
        "order_plan_rows": int(len(order_plan)),
        "history_rows": int(len(order_history)),
        "sheets": {name: int(len(df)) for name, df in sheets.items()},
    }
    write_json(REPORT_PATH, report)
    return report


if __name__ == "__main__":
    print(json.dumps(export_excel_record(), ensure_ascii=False, indent=2))
